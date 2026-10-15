"""Coleta indicadores de TCU (iESGo) e CGU (LAI/Fala.BR e Portal da Transparência)
da UFSCar e grava data/painel.json, consumido pelo painel (static/painel.html).

Uso:  py -m etl.build [--only iesgo,lai,ouv,cor,portal,comp] [--anos 2020-2026]

Fontes (todas oficiais):
  iesgo  - TCU, planilha pública do levantamento iESGo 2024 (não há API)   https://iesgo.tcu.gov.br
  lai    - CGU, dados abertos Fala.BR (pedidos/recursos de acesso à informação)
  portal - CGU, API do Portal da Transparência (exige PORTAL_TRANSPARENCIA_API_KEY)
"""
import argparse
import csv
import io
import json
import os
import statistics as st
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from etl import fonte_cgu

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
CACHE = DATA / "_cache"
OUT = DATA / "painel.json"

IESGO_BASE = "https://iesgo.tcu.gov.br/wp-content/uploads/sites/12/iesgo2024/"
IESGO_TABELA = "iESGo2024_7_TabelaDinamica.xlsx"
IESGO_INDIC = "indicadores-iESGo-2024.xlsx"
LAI_URL, LAI_REC_URL, OUV_URL = fonte_cgu.LAI_URL, fonte_cgu.LAI_REC_URL, fonte_cgu.OUV_URL
PORTAL_URL = "https://api.portaldatransparencia.gov.br/api-de-dados"

UFSCAR_NOME = "Fundação Universidade Federal de São Carlos"
UFSCAR_SIGLA = "UFSCar"


def _load_env() -> None:
    env = ROOT / ".env"
    if env.exists():
        for line in env.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                os.environ.setdefault(k.strip(), v.strip())


_load_env()


def baixar(url: str, destino: Path) -> Path:
    if destino.exists() and destino.stat().st_size > 0:
        return destino
    destino.parent.mkdir(parents=True, exist_ok=True)
    print(f"  baixando {url}", flush=True)
    tmp = destino.with_suffix(destino.suffix + ".part")
    req = urllib.request.Request(url, headers={"User-Agent": "painel-ufscar/1.0"})
    for tentativa in range(4):
        try:
            with urllib.request.urlopen(req, timeout=120) as r, open(tmp, "wb") as f:
                while chunk := r.read(1 << 20):
                    f.write(chunk)
            break
        except OSError:
            if tentativa == 3:
                raise
            time.sleep(5 * (tentativa + 1))
    tmp.replace(destino)
    return destino


# --------------------------------------------------------------------------- TCU / iESGo
FAIXAS = [(0.15, "Inexpressivo"), (0.40, "Iniciando"), (0.70001, "Intermediário"), (1.01, "Aprimorado")]


def faixa(v: float) -> str:
    return next(n for lim, n in FAIXAS if v < lim)


# Rótulos de índices que não constam (ou constam com outra sigla) na planilha de indicadores.
ROTULOS_EXTRA = {
    "iGovSustentAmb": "Governança da sustentabilidade ambiental",
    "iGestSustentAmb": "Gestão da sustentabilidade ambiental",
    "iGestSustentSocial": "Gestão da sustentabilidade social",
    "iGovSustentSocial": "Governança da sustentabilidade social",
    "iGovContratações": "Governança de contratações",
    "iGovPessoas": "Governança de pessoas",
    "iGovTI": "Governança de TI",
    "iGovOrcament": "Governança orçamentária",
    "iGest": "Gestão pública",
    "iES": "Índice de sustentabilidade (ambiental e social)",
}
INDICES = ["iESGo", "iGG", "iES", "iGovPub", "iGest", "Lid", "Estr", "Cont", "iGovPessoas", "iGovTI",
           "iGovContratações", "iGovOrcament", "iGestPessoas", "iGestTI", "iGestContrat", "iGestOrcament",
           "iGovSustentAmb", "iGestSustentAmb", "iGovSustentSocial", "iGestSustentSocial"]


def coletar_iesgo() -> dict:
    import openpyxl
    tabela = baixar(IESGO_BASE + IESGO_TABELA, CACHE / IESGO_TABELA)
    indic = baixar(IESGO_BASE + IESGO_INDIC, CACHE / IESGO_INDIC)

    nomes = dict(ROTULOS_EXTRA)
    ws = openpyxl.load_workbook(indic, read_only=True, data_only=True)["indicadores iESGo2024 x iGG2021"]
    for row in list(ws.iter_rows(values_only=True))[1:]:
        if row[1] and row[2]:
            nomes.setdefault(str(row[1]).strip(), str(row[2]).strip())

    ws = openpyxl.load_workbook(tabela, read_only=True, data_only=True)["iESGo2024"]
    rows = list(ws.iter_rows(values_only=True))
    h = list(rows[0])
    col = {c: i for i, c in enumerate(h) if c}
    faltam = [i for i in INDICES if i not in col]
    if faltam:
        raise RuntimeError(f"Colunas ausentes na planilha do TCU: {faltam}")
    dados = [r for r in rows[1:] if r[col["resp_nome"]]]
    uf = next((r for r in dados if r[col["resp_sigla"]] == UFSCAR_SIGLA), None)
    if uf is None:
        raise RuntimeError("UFSCar não encontrada na planilha iESGo 2024")
    pares = [r for r in dados if r[col["Área.temática"]] == uf[col["Área.temática"]]
             and r[col["Natureza.Jurídica"]] in ("Autarquia", "Fundação")]

    def num(r, i):
        v = r[col[i]]
        return float(v) if isinstance(v, (int, float)) else None

    indices = []
    for ind in INDICES:
        vp = sorted([num(r, ind) for r in pares if num(r, ind) is not None], reverse=True)
        vg = [num(r, ind) for r in dados if num(r, ind) is not None]
        v = num(uf, ind)
        indices.append({
            "sigla": ind, "nome": nomes.get(ind, ind), "ufscar": v, "faixa": faixa(v) if v is not None else None,
            "media_ifes": st.mean(vp), "mediana_ifes": st.median(vp), "min_ifes": min(vp), "max_ifes": max(vp),
            "posicao_ifes": vp.index(v) + 1 if v in vp else None, "n_ifes": len(vp),
            "media_geral": st.mean(vg), "n_geral": len(vg),
        })
    lista = sorted(({"sigla": r[col["resp_sigla"]], "nome": r[col["resp_nome"]], **{k: (round(num(r, k), 4) if num(r, k) is not None else None) for k in INDICES}}
                    for r in pares), key=lambda x: -(x["iESGo"] or 0))
    return {
        "ano": 2024,
        "fonte": "TCU – Levantamento iESGo 2024 (planilha oficial, tabela dinâmica)",
        "url": "https://iesgo.tcu.gov.br",
        "grupo_comparacao": f"{uf[col['Área.temática']]} (autarquias e fundações) – {len(pares)} organizações",
        "faixas": "Inexpressivo <15% · Iniciando 15–40% · Intermediário 40–70% · Aprimorado >70%",
        "indices": indices,
        "pares": lista,
    }


# --------------------------------------------------------------------------- CGU / LAI (Fala.BR)
def _dt(s: str):
    try:
        return datetime.strptime(s.strip(), "%d/%m/%Y").date()
    except ValueError:
        return None


def _eh_ufscar(org: str) -> bool:
    return "UFSCar" in org and "São Carlos" in org


def _eh_universidade(org: str) -> bool:
    return "Universidade Federal" in org or "Universidade Tecnológica Federal" in org


class Acum:
    def __init__(self):
        self.n = 0
        self.dias: list[int] = []
        self.no_prazo = 0
        self.com_prazo = 0
        self.prorrogado = 0
        self.decisao = Counter()
        self.assunto = Counter()
        self.mes = Counter()

    def add_leve(self, r):
        """Só volume e tempo de resposta (usado nos benchmarks)."""
        self.n += 1
        d0, d1, lim = _dt(r["DataRegistro"]), _dt(r["DataResposta"]), _dt(r["PrazoAtendimento"])
        if d0:
            self.mes[d0.month] += 1
        if d0 and d1 and d1 >= d0:
            self.dias.append((d1 - d0).days)
            if lim:
                self.com_prazo += 1
                self.no_prazo += d1 <= lim

    def add(self, r):
        self.add_leve(r)
        self.prorrogado += r["FoiProrrogado"].strip() == "Sim"
        dec = r["Decisao"].strip()
        if dec:
            self.decisao[dec] += 1
        self.assunto[r["AssuntoPedido"].strip() or "Não informado"] += 1

    def resumo(self, detalhado=True):
        out = {
            "pedidos": self.n,
            "tempo_medio_dias": round(st.mean(self.dias), 1) if self.dias else None,
            "tempo_mediano_dias": st.median(self.dias) if self.dias else None,
            "pct_no_prazo": round(100 * self.no_prazo / self.com_prazo, 1) if self.com_prazo else None,
            "pct_prorrogado": round(100 * self.prorrogado / self.n, 1) if self.n else None,
        }
        tot_dec = sum(self.decisao.values())
        concedido = sum(v for k, v in self.decisao.items() if k.startswith("Acesso Concedido") or k.startswith("Acesso Parcialmente"))
        out["pct_acesso_concedido"] = round(100 * concedido / tot_dec, 1) if tot_dec else None
        if detalhado:
            out["decisoes"] = self.decisao.most_common()
            out["assuntos"] = self.assunto.most_common(10)
            out["por_mes"] = [self.mes.get(m, 0) for m in range(1, 13)]
        return out


def _ler_zip(zf: zipfile.ZipFile, nome_csv: str):
    """Lê o CSV AAAAMMDD_<nome_csv> do zip (o prefixo exige o "_", para não pegar SolicitantesPedidos_...)."""
    nome = next(n for n in zf.namelist() if Path(n).name == nome_csv or Path(n).name.endswith("_" + nome_csv))
    csv.field_size_limit(10 ** 9)
    with zf.open(nome) as raw:
        bom = raw.read(2)
    enc = "utf-16" if bom in (b"\xff\xfe", b"\xfe\xff") else "utf-8-sig"
    with zf.open(nome) as raw:
        txt = io.TextIOWrapper(raw, encoding=enc, newline="")
        yield from csv.DictReader(txt, delimiter=";")


def coletar_lai(anos: list[int]) -> dict:
    resultado = []
    for ano in anos:
        zpath, rpath = CACHE / f"lai_{ano}.zip", CACHE / f"lai_rec_{ano}.zip"
        sinc = {}
        for chave, url, destino in ((f"lai_{ano}", LAI_URL, zpath), (f"lai_rec_{ano}", LAI_REC_URL, rpath)):
            try:
                sinc[chave] = fonte_cgu.sincronizar(chave, url.format(ano=ano), destino)
            except Exception as e:  # ano indisponível
                sinc[chave] = {"status": "erro", "erro": str(e)}
        if not zpath.exists():
            print(f"  LAI {ano}: indisponível ({sinc[f'lai_{ano}'].get('erro')})")
            continue
        uf, univ, fed = Acum(), Acum(), Acum()
        rec_inst = Counter()
        with zipfile.ZipFile(zpath) as zf:  # não lê SolicitantesPedidos (perfil dos solicitantes)
            for r in _ler_zip(zf, f"Pedidos_csv_{ano}.csv"):
                org = r["OrgaoDestinatario"]
                if r["Esfera"].strip() == "Federal":
                    fed.add_leve(r)
                if _eh_universidade(org):
                    univ.add_leve(r)
                if _eh_ufscar(org):
                    uf.add(r)
        if rpath.exists():
            with zipfile.ZipFile(rpath) as zf:
                # pelo órgão do pedido: recursos à CGU/CMRI têm outro OrgaoDestinatario
                for r in _ler_zip(zf, f"Recursos_Reclamacoes_csv_{ano}.csv"):
                    if _eh_ufscar(r["OrgaoPedido"]):
                        rec_inst[r["Instancia"].strip() or "—"] += 1
        else:
            print(f"  LAI {ano}: recursos indisponíveis ({sinc[f'lai_rec_{ano}'].get('erro')})")
        item = {"ano": ano, "extracao_cgu": fonte_cgu.data_extracao(zpath), "ufscar": uf.resumo(),
                "universidades_federais": univ.resumo(False), "executivo_federal": fed.resumo(False)}
        item["ufscar"]["recursos"] = sum(rec_inst.values()) if rpath.exists() else None
        item["ufscar"]["recursos_por_instancia"] = rec_inst.most_common()
        # taxa de recursos por 100 pedidos
        if uf.n and rpath.exists():
            item["ufscar"]["recursos_por_100_pedidos"] = round(100 * sum(rec_inst.values()) / uf.n, 1)
        resultado.append(item)
        print(f"  LAI {ano}: {uf.n} pedidos UFSCar | tempo médio {item['ufscar']['tempo_medio_dias']} d", flush=True)
    return {
        "fonte": "CGU – Dados abertos Fala.BR (Lei de Acesso à Informação)",
        "url": "https://buscalai.cgu.gov.br/DownloadDados/DownloadDados",
        "anos": resultado,
        "nota": "Ano corrente é parcial (arquivo atualizado pela CGU periodicamente). Nenhum dado pessoal de solicitantes é armazenado.",
    }


# --------------------------------------------------------------------------- CGU / Ouvidoria (Fala.BR)
def _sem_acento(s: str) -> str:
    import unicodedata
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower().strip()


def coletar_ouvidoria() -> dict:
    """Manifestações de Ouvidoria (e-Ouv/Fala.BR). Guarda linhas anônimas da UFSCar (sem texto nem dados
    demográficos) para filtragem no painel, e médias anuais das universidades federais como referência."""
    zpath = CACHE / "ouvidoria.zip"
    sinc = fonte_cgu.sincronizar("ouvidoria", OUV_URL, zpath)
    if not zpath.exists():
        raise RuntimeError(f"Ouvidoria indisponível: {sinc.get('erro')}")
    csv.field_size_limit(10 ** 9)
    tipos, assuntos, ents = [], [], ["UFSCar", "HU-UFSCar (Ebserh)"]
    ix = lambda lst, v: lst.index(v) if v in lst else (lst.append(v) or len(lst) - 1)  # noqa: E731
    linhas = []
    bench = defaultdict(lambda: {"univ": defaultdict(float), "fed": defaultdict(float)})

    def acum(b, ano, dias, noprazo, sat):
        b["n"] += 1
        if dias is not None:
            b["dias_soma"] += dias
            b["dias_n"] += 1
        if noprazo is not None:
            b["prazo_n"] += 1
            b["prazo_ok"] += noprazo
        if sat is not None:
            b["sat_soma"] += sat
            b["sat_n"] += 1

    with zipfile.ZipFile(zpath) as zf:
        for nome in sorted(zf.namelist()):
            ano = int(nome[-8:-4])
            with zf.open(nome) as raw:
                rd = csv.DictReader(io.TextIOWrapper(raw, encoding="cp1252", newline=""), delimiter=";")
                rd.fieldnames = [_sem_acento(k) for k in rd.fieldnames]  # normaliza o cabeçalho uma vez por arquivo
                for r in rd:
                    org = r["nome orgao"]
                    d0, dp, d1 = _dt(r["data registro"]), _dt(r["data prazo resposta"]), _dt(r["data resposta"])
                    concl = r["situacao"].strip() == "Concluída"
                    dias = (d1 - d0).days if (concl and d0 and d1 and d1 >= d0) else None
                    noprazo = int(d1 <= dp) if (concl and d1 and dp) else None
                    sat = r["satisfacao"].strip()
                    sat = int(sat[1]) if sat[:1] == "(" and sat[1:2].isdigit() else None
                    if "Federal" in org and ("Universidade" in org) and "EBSERH" not in org.upper():
                        acum(bench[ano]["univ"], ano, dias, noprazo, sat)
                    acum(bench[ano]["fed"], ano, dias, noprazo, sat)
                    ent = 0 if "Universidade Federal de São Carlos" in org else (1 if "UFSCar" in org else None)
                    if ent is None:
                        continue
                    at = {"Sim": "S", "Não": "N", "Parcialmente Atendida": "P"}.get(r["demanda atendida"].strip())
                    linhas.append([ano, d0.month if d0 else 0, ix(tipos, r["tipo manifestacao"].strip()), ent,
                                   ix(assuntos, r["assunto"].strip() or "Não informado"),
                                   1 if r["formulario"].strip().startswith("Den") else 0,
                                   dias, noprazo, int(concl), at, sat])

    def resumo(b):
        return {"n": int(b["n"]), "tempo_medio": round(b["dias_soma"] / b["dias_n"], 1) if b["dias_n"] else None,
                "pct_no_prazo": round(100 * b["prazo_ok"] / b["prazo_n"], 1) if b["prazo_n"] else None,
                "satisfacao": round(b["sat_soma"] / b["sat_n"], 2) if b["sat_n"] else None}

    print(f"  Ouvidoria: {len(linhas)} manifestações UFSCar/HU", flush=True)
    return {
        "fonte": "CGU – Dados abertos: manifestações de Ouvidoria (e-Ouv/Fala.BR)",
        "url": "https://www.gov.br/cgu/pt-br/acesso-a-informacao/dados-abertos/arquivos/ouvidoria",
        "colunas": ["ano", "mes", "tipo", "entidade", "assunto", "denuncia", "dias", "no_prazo", "concluida", "atendida", "satisfacao"],
        "tipos": tipos, "assuntos": assuntos, "entidades": ents, "linhas": linhas,
        "referencia": {str(a): {"universidades_federais": resumo(v["univ"]), "executivo_federal": resumo(v["fed"])}
                       for a, v in sorted(bench.items())},
        "nota": "Cada linha é uma manifestação sem texto e sem dados do manifestante. Satisfação e “demanda atendida” são respondidas voluntariamente, por isso cobrem poucas manifestações.",
    }


# --------------------------------------------------------------------------- CGU / Portal da Transparência
class Portal:
    def __init__(self, chave: str):
        self.chave = chave

    def get(self, path: str, **params):
        params = {k: v for k, v in params.items() if v is not None}
        url = f"{PORTAL_URL}{path}?" + urllib.parse.urlencode(params)
        for tentativa in range(5):
            req = urllib.request.Request(url, headers={"chave-api-dados": self.chave, "Accept": "application/json"})
            try:
                with urllib.request.urlopen(req, timeout=60) as r:
                    time.sleep(0.7)  # limite de requisições do Portal
                    return json.load(r)
            except urllib.error.HTTPError as e:
                if e.code == 429 and tentativa < 4:
                    time.sleep(10 * (tentativa + 1))
                    continue
                raise
        return []

    def paginar(self, path: str, max_paginas: int = 400, **params):
        for p in range(1, max_paginas + 1):
            lote = self.get(path, pagina=p, **params)
            if not lote:
                return
            yield from lote
            time.sleep(0.2)


def _brl(s) -> float:
    if isinstance(s, (int, float)):
        return float(s)
    return float(str(s).replace(".", "").replace(",", ".")) if s else 0.0


def coletar_portal(anos: list[int]) -> dict:
    chave = os.getenv("PORTAL_TRANSPARENCIA_API_KEY", "")
    if not chave or "coloque" in chave:
        return {"status": "sem_chave", "mensagem": "Defina PORTAL_TRANSPARENCIA_API_KEY no arquivo .env e rode o ETL novamente."}
    api = Portal(chave)
    siafi = os.getenv("UFSCAR_CODIGO_ORGAO", "26280")
    sup = os.getenv("UFSCAR_ORGAO_SUPERIOR", "26000")
    # valida o código SIAFI antes de agregar qualquer número
    cad = api.get("/orgaos-siafi", codigo=siafi, pagina=1)
    desc = (cad[0].get("descricao", "") if cad else "")
    if "carlos" not in desc.lower():
        return {"status": "erro", "mensagem": f"O código SIAFI {siafi} corresponde a '{desc or 'nada'}', não à UFSCar. Ajuste UFSCAR_CODIGO_ORGAO."}
    out = {"status": "ok", "siafi": siafi, "orgao": desc, "fonte": "CGU – API do Portal da Transparência",
           "url": "https://portaldatransparencia.gov.br/api-de-dados", "despesas": [], "contratos": [], "licitacoes": []}
    for ano in anos:
        d = api.get("/despesas/por-orgao", ano=ano, orgaoSuperior=sup, orgao=siafi, pagina=1)
        d = [x for x in d if str(x.get("codigoOrgao")) == siafi] or d[:1]
        if d:
            out["despesas"].append({"ano": ano, "empenhado": _brl(d[0]["empenhado"]), "liquidado": _brl(d[0]["liquidado"]),
                                    "pago": _brl(d[0]["pago"])})
        ini, fim = f"01/01/{ano}", f"31/12/{ano}"
        ct = list(api.paginar("/contratos", dataInicial=ini, dataFinal=fim, codigoOrgao=siafi))
        out["contratos"].append({"ano": ano, "quantidade": len(ct), "valor_inicial": sum(_brl(c.get("valorInicialCompra")) for c in ct),
                                 "modalidades": Counter(c.get("modalidadeCompra") or "—" for c in ct).most_common(6)})
        lc = list(api.paginar("/licitacoes", dataInicial=ini, dataFinal=fim, codigoOrgao=siafi))
        out["licitacoes"].append({"ano": ano, "quantidade": len(lc), "valor": sum(_brl(c.get("valor")) for c in lc),
                                  "modalidades": Counter(c.get("modalidadeLicitacao") or "—" for c in lc).most_common(6)})
        print(f"  Portal {ano}: {len(ct)} contratos, {len(lc)} licitações", flush=True)
    return out


# --------------------------------------------------------------------------- main
def build_painel(only: str = "iesgo,lai,ouv,cor,comp", anos_str: str | None = None) -> Path:
    if anos_str is None:
        anos_str = f"2020-{date.today().year}"
    ini, fim = (int(x) for x in anos_str.split("-"))
    anos = list(range(ini, fim + 1))
    quais = set(only.split(","))

    atual = json.loads(OUT.read_text(encoding="utf-8")) if OUT.exists() else {}
    painel = {"orgao": {"nome": UFSCAR_NOME, "sigla": UFSCAR_SIGLA}, **{k: atual[k] for k in ("tcu", "cgu") if k in atual}}
    painel.setdefault("tcu", {})
    painel.setdefault("cgu", {})
    if "iesgo" in quais:
        print("TCU iESGo", flush=True)
        painel["tcu"]["iesgo"] = coletar_iesgo()
    if "lai" in quais:
        print("CGU LAI", flush=True)
        painel["cgu"]["lai"] = coletar_lai(anos)
    if "ouv" in quais:
        print("CGU Ouvidoria", flush=True)
        painel["cgu"]["ouvidoria"] = coletar_ouvidoria()
        painel["cgu"]["ouvidoria"]["extracao_cgu"] = fonte_cgu.data_extracao(CACHE / "ouvidoria.zip")
    if "cor" in quais:
        from etl import correicao
        print("CGU Corregedoria (CGU-PAD e CEAF)", flush=True)
        painel["cgu"]["correicao"] = correicao.coletar_correicao()
    if "portal" in quais:
        print("CGU Portal da Transparência", flush=True)
        painel["cgu"]["portal"] = coletar_portal(anos)
    if "comp" in quais:
        from etl import comparativo as cmp
        print("TCU – indicadores das universidades federais (ProPlan + INEP + CAPES)", flush=True)
        painel["tcu"]["indicadores_ufscar"] = cmp.serie_proplan()
        painel["tcu"]["comparativo"] = cmp.comparativo(iesgo=painel["tcu"].get("iesgo"))
    painel["gerado_em"] = datetime.now().isoformat(timespec="seconds")
    DATA.mkdir(exist_ok=True)
    OUT.write_text(json.dumps(painel, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"OK -> {OUT}")
    return OUT


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", default="iesgo,lai,ouv,cor,comp")
    ap.add_argument("--anos", default=f"2020-{date.today().year}")
    a = ap.parse_args(argv)
    build_painel(only=a.only, anos_str=a.anos)


if __name__ == "__main__":
    sys.exit(main())
